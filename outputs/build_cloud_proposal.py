from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT=Path(__file__).parent
doc=Document()
sec=doc.sections[0]
sec.page_height=Inches(11.7); sec.page_width=Inches(8.3)
sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7)
sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=doc.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(26)
doc.styles['Heading 1'].font.size=Pt(18)
doc.styles['Heading 2'].font.size=Pt(12)
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.space_before=Pt(12)
    doc.styles[name].paragraph_format.space_after=Pt(6)
def p(t,style=None): return doc.add_paragraph(t,style)
def h(t): doc.add_heading(t,2)
def page(t): doc.add_page_break(); doc.add_heading(t,1)
def steps(items):
    for i,(title,body) in enumerate(items,1):
        a=p(''); a.add_run(f'{i}. {title} ').bold=True; a.add_run(body)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,txt in zip(t.rows[0].cells,headers): c.text=txt
    for row in rows:
        for c,txt in zip(t.add_row().cells,row): c.text=txt
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr(); no=OxmlElement('w:cantSplit'); pr.append(no)
        if ri==0: pr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]); cp=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            cp.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),'90'); e.set(qn('w:type'),'dxa'); margins.append(e)
            cp.append(margins)
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E9EDF2' if ri==0 else 'FFFFFF'); cp.append(sh)
            va=OxmlElement('w:vAlign'); va.set(qn('w:val'),'center'); cp.append(va)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(3); pp.paragraph_format.line_spacing=1.0
                for r in pp.runs: r.font.size=Pt(10); r.bold=ri==0
    p('')

doc.add_heading('AI Career Service Cloud Migration Proposal',0)
p('Udaan AI team review', 'Subtitle')
p('19 September 2026 | Proposed plan for team finalization')
h('Decision requested')
p('We propose moving the AI Career Service and its supporting development environment to the cloud so teammates can contribute using a browser without downloading large models or running the full application on their computers. The team should finalize the hosting approach, monthly budget, model choices and responsibilities before implementation begins.')
p('The target is cloud development, cloud application hosting, central data storage and cloud AI processing. Student devices will still handle the browser, microphone and audio playback. Reliable internet will be required.')
h('What the current service already provides')
p('The current configuration uses local Ollama with qwen3:1.7b for text and qwen3-embedding:0.6b for embeddings. It also has local English transcription, browser read-aloud, PostgreSQL storage and an internal SearXNG search service. General advisor chat researches the web and requires cited evidence; saved recommendation explanations use the student’s stored results.')
p('Existing features include student-owned question history, refreshed answers and limited topic follow-ups. Recommendation rankings come from programmed assessment rules. Cloud migration should preserve these useful behaviors while replacing local infrastructure and provider connections.')
h('Recommended direction')
p('Use browser-accessible development workspaces, a cloud-hosted website and backend, managed PostgreSQL, private file storage, and a separately configured AI service. Evaluate a hosted text model first; rent a dedicated GPU only if model availability, performance or usage justifies it. Keep development, staging and production data separate.')
h('Migration scope')
p('Move the gateway and required Auth, Student, Assessment and Roadmap dependencies into the cloud environment too. Moving only the AI service would leave key journeys dependent on a teammate’s computer. Reuse the current service structure initially; a complete rewrite is not required to achieve cloud operation.')

page('Cloud architecture and team workflow')
h('Where each part will run')
table(['Component','Proposed cloud responsibility'],[
('Development workspace','Browser editor, dependencies, builds and tests; separate workspace or branch per teammate.'),
('Website and API gateway','Serve the built website over HTTPS and route authenticated requests to private backend services.'),
('Career and supporting services','Run application logic, deterministic scoring, context checks, search orchestration and history.'),
('AI processing','Hosted model API or private GPU endpoint for text; separate embedding and speech connections.'),
('Managed PostgreSQL','Store accounts, profiles, assessments, recommendations and student-owned advisor history.'),
('Private object storage','Store approved documents, exports and backups where needed; self-hosted model files use persistent model storage.'),
('Operations','Manage secrets, access, logs, usage limits, alerts and tested backups.')],[1.7,5.0])
h('Daily work for teammates')
steps([
('Open the workspace.','Sign in to the team’s cloud development environment and open the project in the browser.'),
('Work on a separate change.','Create a branch, edit the relevant service or page, and use test accounts and non-production data.'),
('Verify the change.','Run checks in the cloud workspace and test against the shared development services or an isolated preview.'),
('Request review.','Submit the change for a teammate to review; automated checks build and test the application.'),
('Release through staging.','Merge approved work, verify the staging deployment, then let the assigned release owner promote it to production.')])
p('Teammates will not need local Docker, Python model environments or model downloads for the standard workflow. Cloud workspaces still need persistent storage, access control and automatic shutdown when idle.')

page('Student journey through the cloud service')
steps([
('Sign in and ask.','The student opens Udaan AI, signs in, and types a question or chooses voice input.'),
('Transcribe voice when requested.','The browser records after permission and sends a short audio clip to the backend. Cloud speech processing returns a transcript for the student to review and edit before sending the question. Raw audio is not retained by default.'),
('Validate the request.','The gateway checks authentication and request limits. The Career Service uses the authenticated student identity and applies topic, context and access checks.'),
('Choose the answer route.','General education and career questions use internet research. “Explain my recommendations” reads the student’s current saved assessment and ranking data. Missing or outdated context prompts the student to update it.'),
('Retrieve supporting information.','For general questions, search and read public sources in the cloud, preferring authoritative material and appropriate official sources for admissions or eligibility. Treat website content as evidence, never as instructions.'),
('Generate and validate.','Send only the necessary question and evidence to the text model. Validate response structure and source references. Unsupported questions receive an honest missing-evidence or clarification response; unavailable services receive a retry message.'),
('Show and save the result.','Display the answer, source links and research time. Save completed answers to the student’s own history. Refreshing an old answer triggers new research; history does not become current evidence.'),
('Continue or listen.','A follow-up uses the verified topic and fresh evidence. Read-aloud uses a separately selected speech-output method; the student controls playback.')])
h('Behaviors to preserve')
p('A language model must not silently change assessment scores or recommendation rankings. Interest scores must not be presented as aptitude, guaranteed admission or career-success probabilities. General factual answers require supporting internet sources; citation presence alone does not prove correctness.')
h('Worked example')
p('A student asks, “What can I study after Class 10 for a career in AI?” The cloud service researches relevant education pathways, checks supporting sources and returns a short explanation with links. If the student asks whether they will get selected by a particular college, the service asks for the course and institution and avoids a selection guarantee.')

page('Model choices for team evaluation')
p('The three suggested models perform different jobs. They are candidates for evaluation, not a complete cloud platform or three interchangeable chat models. Model selection should follow measured quality on Udaan AI questions and target languages.')
h('GPT OSS for text generation')
p('The Ollama catalog lists gpt-oss 20B and 120B variants, with local packages of approximately 14 GB and 65 GB, plus cloud variants [1]. Evaluate a hosted 20B option first for cost and response quality. These download sizes are not total running-memory requirements. Confirm endpoint access, structured-output support, concurrency and actual billing before adoption.')
p('A larger text model still needs evidence retrieval and answer validation. Benchmark it against the existing service using the same source material; do not assume that switching models automatically fixes factual errors or latency.')
h('All MiniLM L6 v2 for embeddings')
p('This model converts short text into 384-dimensional vectors for semantic similarity and retrieval. Its model card is English-focused and notes default truncation after 256 word pieces [2]. It is a candidate for an English retrieval baseline, not the answer generator.')
p('If selected, create a new embedding version and rebuild any indexes that depend on the old model. Do not mix old and new vectors. First confirm where embeddings are useful in the current web retrieval pipeline; adding a vector database is not automatically necessary. Evaluate a multilingual alternative before promising Kannada or other Indian-language retrieval.')
h('Shrutam 2 for Indian language transcription')
p('Shrutam-2 is a speech-to-text model for 12 Indian languages, including Kannada and Hindi. Its model card lists a BharatGen non-commercial license and currently shows no hosted Inference Provider [3]. Plan for a separately hosted endpoint if selected, and resolve permitted use before including it in a commercial release.')
p('It does not provide read-aloud. English transcription also needs an explicitly tested solution because English is not in its listed 12-language set. Start with the existing English capability hosted in the cloud or a suitable hosted speech service, then add Indian languages after transcription and answer-quality evaluation.')
h('Proposed first release')
p('Keep English as the initial supported end-to-end language. Move the service to cloud infrastructure, evaluate the text-model replacement, and preserve existing speech behavior where practical. Add Kannada only after speech recognition, retrieval, answer generation and playback have all been tested together.')

page('Implementation plan after team finalization')
steps([
('Record the baseline.','Inventory services, databases, model files, environment settings and current API behavior. Preserve the repository and take a verified backup of existing data. Record a small set of expected student journeys.'),
('Create cloud development and staging.','Set up team accounts, browser workspaces, private service networking, database, storage, secret management and the HTTPS staging website. Deploy supporting services so no developer laptop is required.'),
('Adapt AI connections.','Replace the local-only transport with explicit text, embedding and speech provider interfaces. The current code rejects remote hosts and cloud model names, so changing an address alone is insufficient. Add authenticated HTTPS connections, approved endpoint configuration, timeouts, error handling and server-side secrets.'),
('Move search and storage.','Host SearXNG privately or select an approved search API. Preserve safe page-fetching rules. Apply database migrations and migrate agreed records with ownership intact. Rebuild embeddings only where the selected model requires it.'),
('Prepare for shared use.','Replace process-only capacity controls with coordinated limits or a queue across instances. Separate speech work from text capacity where useful. Retain refresh behavior and cache only impersonal content with appropriate expiry. Measure timeouts under the selected hosting platform.'),
('Validate in staging.','Run automated regression checks and real student journeys. Test evidence quality, model failures, voice permissions, history ownership, concurrency and backup restoration. Record quality, response time and cost per completed interaction.'),
('Run a limited pilot.','Invite a small agreed group. Monitor failures, spending and feedback. Fix defects before wider release; do not promise universal correctness or a response-time improvement before measurement.'),
('Cut over and retain recovery options.','Take a final backup, reconcile changes made during migration, switch the live address and monitor. Keep the prior deployment and backup available for the agreed recovery period. Remove obsolete local dependencies only after cloud acceptance.')])
h('Release and recovery ownership')
p('Assign one owner each for infrastructure and billing, backend and AI, frontend integration, and evaluation and release. One person can hold multiple roles. The release owner controls promotion and recovery; teammates review changes through the shared repository.')
p('If cloud rollout fails, restore the last verified application release and compatible database state using the rehearsed plan. Do not reverse a database migration blindly or delete the existing data to make a deployment succeed.')

page('Decisions required before work starts')
table(['Decision','Team response'],[
('Cloud provider and region','Provider: __________   Region: __________'),
('Development workspaces','Tool: __________   Number of teammates: ____'),
('Monthly budget and billing owner','Budget: __________   Owner: __________'),
('Expected pilot usage','Students: ____   Peak simultaneous requests: ____'),
('Text model and hosting','Model: __________   Hosted API or private GPU: ____'),
('Embedding and search approach','Embedding: __________   Search: __________'),
('Speech and first release languages','Transcription: ______   Playback: ______   Languages: ______'),
('Data handling and recovery','Retention: ______   Backup frequency: ______   Recovery owner: ______'),
('Scope and team agreement','Approved scope: __________   Review date: __________')],[2.4,4.3])
h('Cost and data controls')
p('Estimate monthly cost as development workspaces plus application hosting, database, storage and backups, model usage or GPU runtime, search, speech, network transfer and monitoring. Set usage quotas and alerts, stop idle development resources and define who can authorize increases. A budget alert by itself does not stop spending.')
p('Keep model credentials on the server. Limit staff access, isolate student records and exclude personal content from routine logs. Review provider retention and data-use terms before sending student questions or audio. Do not send full profiles or assessment records to external models unless an approved feature needs them. Document student-facing data handling and deletion behavior.')
h('Acceptance before student release')
p('A teammate can build, test and use the complete service through a browser without local models. Student login, current recommendations, cited chat, follow-ups and private history work in staging. Voice is tested on real devices. Failures return clear messages. The team has approved quality and load results, measured costs, a restored-backup test and the recovery procedure.')
h('Team response')
p('Decision: Approve / Approve with changes / Defer.\nChanges required: __________________________________________________\nImplementation owner: __________________  Agreed start date: ______________')
h('References')
for t in [
'[1] Ollama GPT OSS catalog: https://ollama.com/library/gpt-oss',
'[2] MiniLM model card: https://huggingface.co/sentence-transformers/all-MiniLM-L6-v2',
'[3] Shrutam 2 model card: https://huggingface.co/bharatgenai/Shrutam-2',
'Model pages reviewed in this discussion. Project basis: current AI provider configuration and transport; career answers, local web search and local voice documentation. Vendor, prices and pilot thresholds remain team decisions.'
]:
    pp=p(t)
    for r in pp.runs: r.font.size=Pt(9)
doc.core_properties.title='AI Career Service Cloud Migration Proposal'
doc.core_properties.subject='Team review of cloud workflow and migration decisions'
doc.core_properties.author='Udaan AI'
OUT.mkdir(exist_ok=True)
path=OUT/'AI_Career_Service_Cloud_Migration_Proposal.docx'
doc.save(path)
print(path.resolve())
